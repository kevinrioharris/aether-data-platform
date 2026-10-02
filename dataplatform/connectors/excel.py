"""Read a table out of a human-made spreadsheet.

Spreadsheets are not tables: titles sit above the header, totals below the data, formulas have no
cached values when the file wasn't saved by Excel, and numbers get typed as text. This reader
finds the real table and returns every cell as a string (or None). Typing happens in silver.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date, datetime

from openpyxl import load_workbook


@dataclass(frozen=True)
class ExcelTable:
    columns: list[str]  # normalized snake_case names
    rows: list[tuple[int, list[str | None]]]  # (1-based sheet row number, cell values)


def normalize_column(name: str) -> str:
    return re.sub(r"[^0-9a-z]+", "_", str(name).strip().lower()).strip("_")


def unique_columns(names: list[str], positions: list[int] | None = None) -> list[str]:
    """Make normalized names usable as columns: blank → `column_<position>`, repeats → `<name>_2`, `_3`, ...

    Two headers that normalize alike ("Amount", "amount ") must stay two columns, never overwrite each other.
    """
    positions = positions or list(range(1, len(names) + 1))
    taken: set[str] = set()
    result = []
    for name, pos in zip(names, positions, strict=True):
        base = name or f"column_{pos}"
        candidate, n = base, 1
        while candidate in taken:
            n += 1
            candidate = f"{base}_{n}"
        taken.add(candidate)
        result.append(candidate)
    return result


@dataclass(frozen=True)
class SheetInfo:
    name: str
    state: str  # "visible" | "hidden" | "veryHidden"


def list_sheets(data: bytes) -> list[SheetInfo]:
    wb = load_workbook(io.BytesIO(data), read_only=True)
    try:
        return [SheetInfo(ws.title, ws.sheet_state) for ws in wb.worksheets]
    finally:
        wb.close()


def _to_str(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime | date):
        return value.isoformat()
    text = str(value).strip()
    return text or None


def read_excel_table(
    data: bytes,
    *,
    sheet: str,
    header_row_contains: str | None = None,
    header_row: int | None = None,
    stop_at_first_cell: str | None = None,
    ignore_columns: list[str] | None = None,
) -> ExcelTable:
    """Header row: `header_row` (1-based) if given, else the first row containing `header_row_contains`,
    else the first non-empty row. Columns with an empty header cell are skipped."""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"no sheet {sheet!r} (sheets: {wb.sheetnames})")
    all_rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()

    if header_row is not None:
        header_idx = header_row - 1 if 0 < header_row <= len(all_rows) else None
        missing = f"header row {header_row} is outside the sheet ({len(all_rows)} rows)"
    elif header_row_contains is not None:
        header_idx = next(
            (i for i, row in enumerate(all_rows) if any(_to_str(c) == header_row_contains for c in row)), None
        )
        missing = f"no header row containing {header_row_contains!r}"
    else:
        header_idx = next((i for i, row in enumerate(all_rows) if any(_to_str(c) for c in row)), None)
        missing = "sheet is empty"
    if header_idx is None:
        raise ValueError(f"{missing} in sheet {sheet!r}")

    ignore = {normalize_column(c) for c in ignore_columns or []}
    header = [_to_str(c) for c in all_rows[header_idx]]
    named = [(i, normalize_column(h)) for i, h in enumerate(header) if h and normalize_column(h) not in ignore]
    names = unique_columns([n for _, n in named], positions=[i + 1 for i, _ in named])
    keep = list(zip([i for i, _ in named], names, strict=True))

    rows = []
    for offset, row in enumerate(all_rows[header_idx + 1 :], start=header_idx + 2):
        values = [_to_str(row[i]) if i < len(row) else None for i, _ in keep]
        if stop_at_first_cell and values and values[0] == stop_at_first_cell:
            break
        if any(v is not None for v in values):
            rows.append((offset, values))
    return ExcelTable(columns=[name for _, name in keep], rows=rows)
