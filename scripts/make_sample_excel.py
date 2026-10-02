"""Create a realistic (i.e. slightly messy) Sales Ops target workbook in sample_data/.

    python scripts/make_sample_excel.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

# Monthly revenue targets in USD. Display names differ from the product DB on purpose.
TARGETS_2026 = {
    "Books": [22000, 20000, 24000, 24000, 25000, 26000, 30000, 28000, 40000, 36000, 42000, 55000],
    "Electronics": [30000, 28000, 32000, 32000, 33000, 34000, 34000, 32000, 50000, 45000, 60000, 80000],
    "Fashion": [20000, 19000, 23000, 24000, 24000, 25000, 25000, 26000, 35000, 33000, 40000, 50000],
    "Home & Living": [24000, 22000, 26000, 27000, 28000, 29000, 30000, 25000, 38000, 36000, 42000, 52000],
    "Sports": [30000, 28000, 34000, 36000, 38000, 40000, 40000, 38000, 60000, 50000, 48000, 55000],
}


def build(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Targets"
    ws["A1"] = "FY2026 Revenue Targets by Category (USD)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Owner: Sales Ops · Approved by CFO · last updated 2026-06-15"

    header_row = 4
    for col, name in enumerate(["Category", *MONTHS, "FY Total"], start=1):
        ws.cell(row=header_row, column=col, value=name).font = Font(bold=True)

    for i, (category, values) in enumerate(TARGETS_2026.items()):
        row = header_row + 1 + i
        ws.cell(row=row, column=1, value=category)
        for m, value in enumerate(values, start=2):
            ws.cell(row=row, column=m, value=value)
        ws.cell(row=row, column=14, value=f"=SUM(B{row}:M{row})")

    # Someone typed one number as text, with a thousands separator.
    ws.cell(row=header_row + 2, column=9, value="32,000")  # Electronics / Aug

    total_row = header_row + 1 + len(TARGETS_2026)
    ws.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
    for col in range(2, 15):
        letter = ws.cell(row=header_row, column=col).column_letter
        ws.cell(row=total_row, column=col, value=f"=SUM({letter}{header_row + 1}:{letter}{total_row - 1})")

    notes = wb.create_sheet("Notes")
    notes["A1"] = "Targets are net revenue in USD after cancellations."
    notes["A2"] = "Q4 includes holiday campaign uplift."

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


if __name__ == "__main__":
    out = Path(__file__).resolve().parent.parent / "sample_data" / "sales_targets_2026.xlsx"
    build(out)
    print(f"wrote {out}")
